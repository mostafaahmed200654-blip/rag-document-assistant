"""
Document ingestion utilities: text extraction from PDF / DOCX / TXT / MD
files, and a simple, dependency-light sliding-window chunker.
"""

import os
import re
from typing import List

from pypdf import PdfReader
from docx import Document as DocxDocument

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileTypeError(Exception):
    pass


def get_file_extension(filename: str) -> str:
    return os.path.splitext(filename)[1].lower()


def extract_text_from_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages_text = []
    for page_number, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages_text.append(text)
    return "\n\n".join(pages_text)


def extract_text_from_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of any tables in the document.
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n\n".join(paragraphs)


def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_text(file_path: str, filename: str) -> str:
    """Dispatch to the correct extractor based on file extension."""
    extension = get_file_extension(filename)

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)
    elif extension == ".docx":
        return extract_text_from_docx(file_path)
    elif extension in (".txt", ".md"):
        return extract_text_from_txt(file_path)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension}'. Supported types: {sorted(SUPPORTED_EXTENSIONS)}"
        )


def clean_text(text: str) -> str:
    """Normalize whitespace without destroying paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, chunk_size: int = 800, chunk_overlap: int = 120) -> List[str]:
    """
    Split text into overlapping, word-based chunks.

    A word-based sliding window is used (rather than naive character
    slicing) so that chunk boundaries do not cut words in half, which
    keeps embeddings and retrieved context coherent.

    Args:
        text: The full document text.
        chunk_size: Target number of words per chunk.
        chunk_overlap: Number of words shared between consecutive chunks.

    Returns:
        A list of text chunks in original document order.
    """
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    text = clean_text(text)
    words = text.split()

    if not words:
        return []

    chunks = []
    step = chunk_size - chunk_overlap
    start = 0

    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk_words = words[start:end]
        chunk = " ".join(chunk_words).strip()
        if chunk:
            chunks.append(chunk)
        if end == len(words):
            break
        start += step

    return chunks


def process_document(file_path: str, filename: str, chunk_size: int, chunk_overlap: int) -> List[str]:
    """Full pipeline: extract raw text from a file, then chunk it."""
    raw_text = extract_text(file_path, filename)
    if not raw_text or not raw_text.strip():
        raise ValueError(f"No extractable text found in '{filename}'. The file may be empty, scanned, or corrupted.")
    return chunk_text(raw_text, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
